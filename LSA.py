import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin
import time
import os
from openpyxl import Workbook, load_workbook

def get_product_links(total_pages):
    base_url = "https://www.eldorado-international.com"
    page_template = "https://www.eldorado-international.com/l-s-a/?page={}"
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }
    
    all_links = []
    print("--- PHASE 1: GATHERING PRODUCT LINKS ---")
    
    for page in range(1, total_pages + 1):
        print(f"Scanning Page {page} of {total_pages}...")
        url = page_template.format(page)
        
        try:
            response = requests.get(url, headers=headers)
            response.raise_for_status()
            soup = BeautifulSoup(response.text, 'html.parser')
            
            product_containers = soup.find_all('div', class_='product-container')
            for container in product_containers:
                a_tag = container.find('a')
                if a_tag and 'href' in a_tag.attrs:
                    full_link = urljoin(base_url, a_tag['href'])
                    all_links.append(full_link)
            
            time.sleep(1) # Brief pause between pagination pages
            
        except requests.exceptions.RequestException as e:
            print(f"Error on page {page}: {e}")

    print(f"\nTotal links scraped: {len(all_links)}\n")
    return all_links

def scrape_product_details(product_url, headers):
    try:
        response = requests.get(product_url, headers=headers)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, 'html.parser')
        
        product_data = {}

        # 1. Extract Table Details
        table_rows = soup.select('.tab-content table tr')
        for row in table_rows:
            columns = row.find_all('td')
            if len(columns) == 2:
                category = columns[0].text.replace('\xa0', '').strip()
                value = columns[1].text.replace('\xa0', '').strip()
                if category.endswith(':'):
                    category = category[:-1].strip()
                product_data[category] = value

        # 2. Extract Image URL
        image_element = soup.select_one('#product_image_swiper .swiper-slide-inside img')
        if image_element and 'src' in image_element.attrs:
            base_url = "https://www.eldorado-international.com"
            product_data['Image URL'] = urljoin(base_url, image_element['src'])
        else:
            product_data['Image URL'] = ""

        return product_data

    except requests.exceptions.RequestException as e:
        print(f"Error scraping product details: {e}")
        return None

def main(total_pages):
    excel_filename = "LSA.xlsx"
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }

    # Set up the Excel Workbook
    if not os.path.exists(excel_filename):
        wb = Workbook()
        ws = wb.active
        ws.title = "Scraped Data"
        # Setup Column Headers (Optional, but helps with organization)
        ws.append(["Blank_A", "Blank_B", "SKU", "Brand", "Blank_E", "Image URL", "EAN", "Name", "Outer Carton", "Color"])
        wb.save(excel_filename)
    else:
        wb = load_workbook(excel_filename)
        ws = wb.active

    # Phase 1: Get all links
    product_links = get_product_links(total_pages)
    
    if not product_links:
        print("No links found to scrape. Exiting.")
        return

    # Phase 2: Scrape details and save to Excel
    print("--- PHASE 2: SCRAPING DETAILS & SAVING ---")
    
    for index, link in enumerate(product_links, start=1):
        print(f"\nProcessing [{index}/{len(product_links)}]: {link}")
        
        details = scrape_product_details(link, headers)
        
        if details:
            # Print to terminal
            for key, val in details.items():
                print(f"  {key}: {val}")

            # Map the dictionary to specific Excel columns
            # List indices map to Excel columns: 0=A, 1=B, 2=C, etc.
            row_data = [
                "",                              # A (Blank)
                "",                              # B (Blank)
                details.get("SKU", ""),          # C
                details.get("Brand", ""),        # D
                "",                              # E (Blank)
                details.get("Image URL", ""),    # F
                details.get("EAN", ""),          # G
                details.get("Name", ""),         # H
                details.get("Outer Carton", ""), # I
                details.get("Color", "")         # J
            ]
            
            # Append the row to the worksheet
            ws.append(row_data)

        # Save to Excel every 5 links to avoid complication/data loss
        if index % 5 == 0:
            wb.save(excel_filename)
            print(f"  [>>> Checkpoint: Saved up to item {index} in {excel_filename} <<<]")
        
        # Polite delay to avoid hammering the server
        time.sleep(1.5)

    # Final save to catch any remaining items that didn't hit the mod 5 condition
    wb.save(excel_filename)
    print(f"\n=== COMPLETE! All data saved to {excel_filename} ===")

if __name__ == "__main__":
    # For testing, you can change this to 1 or 2. 
    # Change to 31 when you want to run the full scrape.
    TEST_PAGES = 31 
    main(total_pages=TEST_PAGES)